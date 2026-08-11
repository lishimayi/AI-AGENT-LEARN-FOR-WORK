"""文档处理模块 - 解析、切片、索引"""
import os
import re
from pathlib import Path
from typing import List, Dict, Optional, Any
from abc import ABC, abstractmethod
from dataclasses import dataclass

from bs4 import BeautifulSoup
import html2text
from langchain.text_splitter import RecursiveCharacterTextSplitter
from loguru import logger

from src.core.config import config


@dataclass
class DocumentChunk:
    """文档块"""
    content: str
    metadata: Dict[str, Any]
    chunk_id: str
    word_count: int


class BaseParser(ABC):
    """文档解析器基类"""
    
    @property
    @abstractmethod
    def supported_extensions(self) -> List[str]:
        """支持的扩展名"""
        pass
    
    @abstractmethod
    def parse(self, file_path: str) -> str:
        """解析文档"""
        pass


class PDFParser(BaseParser):
    """PDF 解析器"""
    
    @property
    def supported_extensions(self) -> List[str]:
        return [".pdf"]
    
    def parse(self, file_path: str) -> str:
        from PyPDF2 import PdfReader
        
        text_parts = []
        reader = PdfReader(file_path)
        
        for page_num, page in enumerate(reader.pages, 1):
            text = page.extract_text()
            if text:
                text_parts.append(f"[第 {page_num} 页]\n{text}")
        
        return "\n\n".join(text_parts)


class DocxParser(BaseParser):
    """Word 文档解析器"""
    
    @property
    def supported_extensions(self) -> List[str]:
        return [".docx"]
    
    def parse(self, file_path: str) -> str:
        from docx import Document
        
        doc = Document(file_path)
        text_parts = []
        
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)
        
        return "\n\n".join(text_parts)


class TextParser(BaseParser):
    """纯文本解析器"""
    
    @property
    def supported_extensions(self) -> List[str]:
        return [".txt", ".md", ".markdown"]
    
    def parse(self, file_path: str) -> str:
        encodings = ["utf-8", "gbk", "gb2312", "iso-8859-1"]
        
        for encoding in encodings:
            try:
                with open(file_path, "r", encoding=encoding) as f:
                    return f.read()
            except UnicodeDecodeError:
                continue
        
        raise ValueError(f"无法解析文件: {file_path}")


class HTMLParser(BaseParser):
    """HTML 解析器"""
    
    @property
    def supported_extensions(self) -> List[str]:
        return [".html", ".htm"]
    
    def parse(self, file_path: str) -> str:
        with open(file_path, "r", encoding="utf-8") as f:
            html_content = f.read()
        
        h = html2text.HTML2Text()
        h.ignore_links = False
        h.ignore_images = True
        
        return h.handle(html_content)


class PPTXParser(BaseParser):
    """PowerPoint 解析器"""
    
    @property
    def supported_extensions(self) -> List[str]:
        return [".pptx"]
    
    def parse(self, file_path: str) -> str:
        try:
            from pptx import Presentation
            
            prs = Presentation(file_path)
            text_parts = []
            
            for slide_num, slide in enumerate(prs.slides, 1):
                slide_text = []
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        slide_text.append(shape.text)
                
                if slide_text:
                    text_parts.append(f"[幻灯片 {slide_num}]\n" + "\n".join(slide_text))
            
            return "\n\n".join(text_parts)
        except ImportError:
            raise ImportError("需要安装 python-pptx 库来解析 PPT 文件")


class DocumentProcessor:
    """文档处理器"""
    
    def __init__(self, doc_config: Optional[Dict] = None):
        self.config = doc_config or config.document
        
        # 初始化解析器
        self.parsers: Dict[str, BaseParser] = {}
        for parser in [PDFParser(), DocxParser(), TextParser(), HTMLParser(), PPTXParser()]:
            for ext in parser.supported_extensions:
                self.parsers[ext] = parser
        
        # 文本分割器
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=self.config.get("chunk_size", 500),
            chunk_overlap=self.config.get("chunk_overlap", 50),
            length_function=len,
            separators=["\n\n", "\n", "。", "！", "？", " ", ""]
        )
        
        logger.info(f"DocumentProcessor initialized with {len(self.parsers)} parsers")
    
    def get_parser(self, file_path: str) -> Optional[BaseParser]:
        """获取文件对应的解析器"""
        ext = Path(file_path).suffix.lower()
        return self.parsers.get(ext)
    
    def parse_document(self, file_path: str) -> str:
        """解析文档"""
        parser = self.get_parser(file_path)
        if not parser:
            raise ValueError(f"不支持的文件格式: {file_path}")
        
        logger.info(f"Parsing document: {file_path}")
        return parser.parse(file_path)
    
    def chunk_text(self, text: str, metadata: Dict) -> List[DocumentChunk]:
        """将文本分块"""
        chunks = self.text_splitter.split_text(text)
        
        result = []
        for i, chunk in enumerate(chunks):
            chunk_id = self._generate_chunk_id(chunk, i)
            
            chunk_metadata = {
                **metadata,
                "chunk_index": i,
                "total_chunks": len(chunks)
            }
            
            result.append(DocumentChunk(
                content=chunk,
                metadata=chunk_metadata,
                chunk_id=chunk_id,
                word_count=len(chunk)
            ))
        
        logger.info(f"Created {len(result)} chunks from text")
        return result
    
    def _generate_chunk_id(self, content: str, index: int) -> str:
        """生成 chunk ID"""
        import hashlib
        content_hash = hashlib.md5(content.encode()).hexdigest()[:8]
        return f"chunk_{index}_{content_hash}"
    
    def process_file(self, file_path: str) -> List[DocumentChunk]:
        """处理单个文件"""
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"文件不存在: {file_path}")
        
        # 解析文档
        text = self.parse_document(file_path)
        
        # 获取文件元数据
        file_stats = os.stat(file_path)
        path = Path(file_path)
        
        metadata = {
            "source": str(path.absolute()),
            "file_name": path.name,
            "file_type": path.suffix.lower(),
            "file_size": file_stats.st_size,
            "created_time": file_stats.st_ctime,
            "modified_time": file_stats.st_mtime,
        }
        
        # 自动识别文档类型
        metadata["doc_type"] = self._infer_doc_type(path.name, text)
        
        # 分块
        return self.chunk_text(text, metadata)
    
    def _infer_doc_type(self, filename: str, content: str) -> str:
        """推断文档类型"""
        filename_lower = filename.lower()
        content_lower = content.lower()
        
        if "prd" in filename_lower or "产品需求" in content_lower:
            return "prd"
        elif "api" in filename_lower or "接口" in content_lower:
            return "interface"
        elif "流程" in content_lower or "flow" in filename_lower:
            return "workflow"
        elif "架构" in content_lower or "架构" in filename_lower:
            return "architecture"
        elif "测试" in content_lower:
            return "test"
        else:
            return "general"
    
    def process_batch(self, file_paths: List[str]) -> List[DocumentChunk]:
        """批量处理文件"""
        all_chunks = []
        
        for file_path in file_paths:
            try:
                chunks = self.process_file(file_path)
                all_chunks.extend(chunks)
            except Exception as e:
                logger.error(f"处理文件失败 {file_path}: {e}")
        
        logger.info(f"Batch processed {len(file_paths)} files, created {len(all_chunks)} chunks")
        return all_chunks
    
    def chunks_to_dict(self, chunks: List[DocumentChunk]) -> List[Dict]:
        """将 DocumentChunk 转换为字典格式"""
        return [
            {
                "content": chunk.content,
                "metadata": chunk.metadata,
                "chunk_id": chunk.chunk_id,
                "word_count": chunk.word_count
            }
            for chunk in chunks
        ]
